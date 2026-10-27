"""Lossless workbook support shared by the ANNX and PCR template parsers.

Source tables retain coordinates, cached values and formulas. Scientific values
are never recalculated: openpyxl is not an Excel calculation engine.
"""
import csv
import logging
import math
import re
from dataclasses import asdict, dataclass
from datetime import date, datetime
from typing import Any

import openpyxl
from openpyxl.utils.datetime import from_excel

logger = logging.getLogger(__name__)


def normalize_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


@dataclass
class Measurement:
    endpoint: str
    treatment: str
    value: Any
    unit: str
    source_cell: str
    replication: str | None = None
    sd: Any = None
    gate: str | None = None


class CellAssayParser:
    test_type = ""

    def __init__(self, file_path: str):
        self.warnings: list[str] = []
        self.wb = openpyxl.load_workbook(file_path, data_only=True)
        try:
            self.formulas = openpyxl.load_workbook(file_path, data_only=False)
            matches = [s for s in self.wb if normalize_key(s.title) == "test_conditions"]
            if len(matches) != 1:
                raise ValueError("Expected one 'Test conditions' or 'Test_conditions' sheet.")
            self.ws = matches[0]
            acronym = self.lookup("short_name_or_acronym")
            if normalize_key(acronym) != normalize_key(self.test_type):
                raise ValueError(f"Expected {self.test_type} workbook; test acronym is {acronym!r}.")
        except Exception:
            self.close()
            raise

    def close(self):
        self.wb.close()
        if hasattr(self, "formulas"):
            self.formulas.close()

    def warn(self, message: str):
        if message not in self.warnings:
            self.warnings.append(message)
            logger.warning("%s: %s", self.test_type, message)

    def value(self, value: Any):
        if isinstance(value, (date, datetime)):
            return value.strftime("%Y-%m-%d")
        if isinstance(value, float) and not math.isfinite(value):
            return None
        return value.strip() if isinstance(value, str) else value

    def normalized_date(self, value: Any, location: str):
        if value is None:
            return None
        if isinstance(value, (date, datetime)):
            return value.strftime("%Y-%m-%d")
        if isinstance(value, (int, float)):
            try:
                return from_excel(value, self.wb.epoch).strftime("%Y-%m-%d")
            except (ValueError, OverflowError, AttributeError):
                pass
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y"):
            try:
                return datetime.strptime(str(value).strip(), fmt).strftime("%Y-%m-%d")
            except ValueError:
                pass
        self.warn(f"Malformed date at {location}: {value!r}; retained unchanged.")
        return self.value(value)

    def lookup(self, prefix: str):
        for row in self.ws:
            if normalize_key(row[0].value).startswith(prefix):
                return self.value(row[1].value)
        return None

    def sheet(self, name: str):
        return next((s for s in self.wb if normalize_key(s.title) == normalize_key(name)), None)

    def table(self, ws):
        """Sparse source rows; ignore formatting-only trailing rows, not measurements."""
        rows = []
        formulas = self.formulas[ws.title]
        for row in formulas:
            cells = []
            for cell in row:
                if cell.value is None or (isinstance(cell.value, str) and not cell.value.strip()):
                    continue
                cached = ws[cell.coordinate]
                entry = {"column": cell.column_letter, "coordinate": cell.coordinate,
                         "value": self.value(cached.value)}
                if cell.data_type == "f":
                    entry["formula"] = cell.value
                    if cached.value is None:
                        self.warn(f"Missing cached formula value at {ws.title}!{cell.coordinate}; open and recalculate in Excel.")
                if cell.data_type == "e" or cached.data_type == "e":
                    self.warn(f"Excel error at {ws.title}!{cell.coordinate}: {cached.value}")
                if isinstance(cell.value, str) and "â„" in cell.value:
                    self.warn(f"Possible text encoding issue in {ws.title}; original labels retained.")
                cells.append(entry)
            if cells:
                rows.append({"row": row[0].row, "cells": cells})
        return {"sheet_name": ws.title, "rows": rows,
                "merged_ranges": [str(r) for r in ws.merged_cells.ranges]}

    def extract_details(self):
        wp_fields = {"wp_name": "project_work_package", "partner": "partner_conducting",
                     "laboratory_name": "test_facility", "full_test_name": "full_name_of_test",
                     "test_acronym": "short_name_or_acronym", "test_type": "type_or_class",
                     "endpoint": "end_point", "endpoint_outcome": "metric_s", "sop": "sop_s", "path": "path_link"}
        mat_fields = {"material_identifier": "sample_cms", "erm_id": "erm_identifier",
                      "material_name": "material_name", "core_chemistry": "core_chemistry",
                      "cas_no": "cas_no", "cas_for_core": "cas_for_core", "material_supplier": "material_supplier",
                      "material_state": "material_state", "batch": "batch", "vial": "vial"}
        wp = {key: self.lookup(prefix) for key, prefix in wp_fields.items()}
        for role, prefix in [("lead_scientists", "lead_scientist"), ("assay_scientists", "assay_test_work")]:
            scientists = []
            for row in self.ws:
                if normalize_key(row[0].value).startswith(prefix) or (
                    role == "assay_scientists" and not row[0].value and row[1].value
                    and "mail" in str(row[2].value).lower()
                ):
                    scientists.append({"name": self.value(row[1].value), "email": self.value(row[3].value)})
            wp[role] = scientists
        material = {key: self.lookup(prefix) for key, prefix in mat_fields.items()}
        material["preparation_date"] = self.normalized_date(self.lookup("date_of_preparation"), "preparation date")
        replications = []
        replication_header_row = None
        for row in self.ws:
            identifier = row[1].value
            if re.fullmatch(r"WP\d+_(?:ANX|ANNX|PCR)_.+R\d+", str(identifier or "")):
                if replication_header_row is None:
                    replication_header_row = row[0].row - 1
                item = {"test_identifier": identifier}
                for cell in row[2:]:
                    if cell.value is None:
                        continue
                    header = normalize_key(self.ws.cell(replication_header_row, cell.column).value)
                    if "date" in header:
                        item[header] = self.normalized_date(cell.value, cell.coordinate)
                    elif isinstance(cell.value, str) and "replicate" in cell.value.lower():
                        item["replicate_label"] = cell.value
                replications.append(item)
        # The complete source table retains secondary endpoints, multi-column
        # treatments, passage numbers, units and assay-specific conditions.
        sections = []
        current = {"title": "Test information", "fields": []}
        for row in self.ws:
            label = row[0].value
            if not isinstance(label, str) or not label.strip():
                continue
            values = [self.value(c.value) for c in row[1:]]
            while values and values[-1] is None:
                values.pop()
            if label.strip().isupper() and len(label.strip()) > 5:
                if current["fields"]:
                    sections.append(current)
                current = {"title": label.strip(), "fields": []}
            if values:
                current["fields"].append({"key": normalize_key(label), "label": label.strip(), "values": values})
            if "oncentration" in label and "concentration" not in label.lower():
                self.warn(f"Spelling issue at {self.ws.title}!A{row[0].row}: {label.strip()}")
        if current["fields"]:
            sections.append(current)
        return {"work_package": wp, "material": material, "replications": replications,
                "experimental_sections": sections, "source": self.table(self.ws), "parser_warnings": self.warnings}

    def replication_id(self, title: str):
        match = re.search(r"WP\d+_(?:ANX|ANNX|PCR)_.+R\d+", title)
        return match.group() if match else title

    def raw_blocks(self):
        raise NotImplementedError

    def processed_blocks(self):
        raise NotImplementedError

    def final_measurements(self, ws):
        raise NotImplementedError

    def parse_all_data(self):
        try:
            details = self.extract_details()
            raw = self.raw_blocks()
            processed = self.processed_blocks()
            final_ws = self.sheet("Final results")
            final = {"measurements": self.final_measurements(final_ws), "source": self.table(final_ws)} if final_ws else {}
            if not final_ws:
                self.warn("Optional Final results sheet is absent.")
            if not raw:
                raise ValueError(f"No {self.test_type} raw/output replicate sheets found.")
            for meta in details["replications"]:
                identifier = meta["test_identifier"]
                if not any(b["replication"] == identifier for b in raw):
                    self.warn(f"Missing raw/output sheet for {identifier}.")
                if not any(b["replication"] == identifier for b in processed):
                    self.warn(f"Missing processed sheet for {identifier}.")
            result = {"test_details": details, "replications": raw,
                      "processed_data": {"blocks": processed}, "final_results": final}
            stats = self.sheet("Statistics") or self.sheet("Statistical analysis")
            if stats:
                result["statistical_analysis"] = {"available": True, "source": self.table(stats)}
            logger.info("Parsed %s: %d raw and %d processed replicates", self.test_type, len(raw), len(processed))
            return result
        finally:
            self.close()

    def measurement(self, ws, row, col, endpoint, treatment, unit, **kwargs):
        return asdict(Measurement(str(endpoint).strip(), str(treatment).strip(),
                                 self.value(ws.cell(row, col).value), unit,
                                 f"{ws.title}!{ws.cell(row, col).coordinate}", **kwargs))

    def flow_records(self, ws):
        """Read both CSV-in-cell exports and their split-column counterparts."""
        headers = None
        records = []
        for row in ws:
            values = [self.value(c.value) for c in row]
            if isinstance(values[0], str) and values[0].startswith("Plate,"):
                headers = next(csv.reader([values[0]]))
                continue
            if values[0] == "Plate":
                headers = [str(v) for v in values if v is not None]
                continue
            if not headers or not isinstance(values[0], str) or not values[0].startswith("CMS"):
                continue
            if ",Experiment," in values[0]:
                values = next(csv.reader([values[0]]))
                values = [self.flow_value(v) for v in values]
            if len(values) < len(headers):
                self.warn(f"Truncated cytometry row at {ws.title}!{row[0].row}.")
            keys = [normalize_key(h.replace("%", " percent ")) for h in headers]
            if len(set(keys)) != len(keys):
                self.warn(f"Ambiguous cytometry headers in {ws.title}; columns retained in source table.")
            records.append({"source_row": row[0].row, **dict(zip(keys, values))})
        return records

    @staticmethod
    def flow_value(value):
        if value in ("", "N/A"):
            return value or None
        try:
            return float(value)
        except ValueError:
            return value