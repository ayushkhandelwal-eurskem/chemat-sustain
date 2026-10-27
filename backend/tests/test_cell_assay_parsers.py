"""Self-contained templates: no dependency on users' private Downloads files."""
import json
from datetime import datetime

import pytest
from openpyxl import Workbook

from parsers.annx import parse_excel_annx
from parsers.pcr import parse_excel_pcr


def make_workbook(path, test_type="PCR", *, statistics=True):
    wb = Workbook()
    ws = wb.active
    ws.title = "Test_conditions" if test_type == "PCR" else "Test conditions"
    for row, values in {
        11: ["Project Work Package:", "WP3"],
        14: ["Lead Scientist & contact for test:", "Lead", "E-mail address:", "lead@example.org"],
        15: ["Assay/Test work conducted by:", "Scientist", "E-mail address:", "one@example.org"],
        16: [None, "Second scientist", "E-mail address:", "two@example.org"],
        18: ["Short name or acronym for test/assay:", test_type],
        25: ["Sample CMS Internal Identifier", "CMS_1a_AuNP"],
        35: ["Date of preparation:", datetime(2025, 2, 13)],
        39: [None, None, "Test start date", "Test End Date", "No of replicate"],
        40: ["Test Identifier Number:", f"WP3_{'ANX' if test_type == 'ANNX' else 'PCR'}_1aR1", datetime(2026, 2, 16), "bad date", "Replicate 1"],
        75: ["Treatment concentration series unit:", "μg/mL"],
        76: ["Treatment concentration series labels", "CTRL", "VEHICLE", "EC10", "EC25"],
        77: ["Treatment concentration series", 0, 14, 0.46, 3.03],
    }.items():
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
    identifier = f"WP3_{'ANX' if test_type == 'ANNX' else 'PCR'}_1aR1"
    if test_type == "PCR":
        raw = wb.create_sheet(f"Raw_data_{identifier}")
        raw.append([identifier, None, "Mean Ct", "Sample 1", "Sample 2", "Sample 3", "Sample 4"])
        raw.append([None, None, "GAPDH", 10, 11, 12, 13])
        raw.append([None, None, "IL-6", 20, 21, None, 23])
        processed = wb.create_sheet(f"Processed data_{identifier}")
        processed.append(["# 2^-ΔΔCt"])
        processed.append(["Fold change vs CTRL", "CTRL", "VEHICLE", "EC10", "EC25"])
        processed.append(["IL-6", 1, 2, 3, 4])
        final = wb.create_sheet("Final results")
        final.append([None, "IL-6 (Fold change vs CTRL)", "CTRL", "VEHICLE", "EC10", "EC25"])
        final.append([None, identifier, 1, 2, 3, 4])
        final.append([None, "MEAN", 1, 2, 3, 4])
        final.append([None, "SD", 0, 0.2, 0.3, 0.4])
    else:
        raw = wb.create_sheet(f"Output data_{identifier}")
        raw.append(["Plate,Experiment,Group,Filename,Sample,Workspace,Compensation Source,Plot Title,Gate,Autogate Status,X Parameter,Y Parameter,Count,Volume,Concentration,%Total,%Gated"])
        raw.append(["CMS - AuNP,Experiment,Group(1),nc.fcs,NC,Experiment,Experiment,,AnnV+/DAPI-,N/A,X,Y,200,94,2.1,1.2,2.74"])
        processed = wb.create_sheet(f"Processed data_{identifier}")
        processed.append(raw.cell(1, 1).value.split(","))
        processed.append(["CMS - AuNP", "Experiment", "Group(1)", "nc.fcs", "NC", "Experiment", "Experiment", None, "AnnV+/DAPI-", "N/A", "X", "Y", 200, 94, 2.1, 1.2, 2.74])
        final = wb.create_sheet("Final results")
        final.cell(9, 3, identifier)
        final.cell(10, 5, "NC")
        for col, value in enumerate(["Early apoptotic", "AnnV+/DAPI-", 2.74], 3):
            final.cell(11, col, value)
        for col, value in enumerate(["Early apoptotic", "AnnV+/DAPI-", "Mean", 2.74], 2):
            final.cell(57, col, value)
        final.cell(58, 5, 0.2)
    # Formatting-only dimensions must never result in empty output rows.
    raw.cell(1000, 26).number_format = "0.00"
    if statistics:
        stats = wb.create_sheet("Statistics" if test_type == "PCR" else "Statistical analysis")
        stats.append(["ANOVA summary", "IL-6"])
        stats.append(["P value", 0.0285])
        stats.append(["Tukey", "CTRL vs. EC10", "ns"])
    wb.save(path)
    wb.close()
    return path


@pytest.mark.parametrize("test_type,parser", [("PCR", parse_excel_pcr), ("ANNX", parse_excel_annx)])
def test_assay_contract_metadata_sources_and_statistics(tmp_path, test_type, parser):
    data = parser(str(make_workbook(tmp_path / "assay.xlsx", test_type)))
    json.dumps(data, allow_nan=False)
    assert set(data) == {"test_details", "replications", "processed_data", "final_results", "statistical_analysis"}
    details = data["test_details"]
    assert details["material"]["preparation_date"] == "2025-02-13"
    assert details["work_package"]["wp_name"] == "WP3"
    assert len(details["work_package"]["assay_scientists"]) == 2
    assert details["replications"][0]["test_start_date"] == "2026-02-16"
    assert any("Malformed date" in warning for warning in details["parser_warnings"])
    assert len(data["replications"][0]["source"]["rows"]) < 10
    assert data["statistical_analysis"]["source"]["rows"][1]["cells"][1]["value"] == 0.0285
    assert data["final_results"]["measurements"]


def test_pcr_preserves_missing_values_and_reported_sd(tmp_path):
    data = parse_excel_pcr(str(make_workbook(tmp_path / "pcr.xlsx")))
    raw = data["replications"][0]["measurements"]
    assert len(raw) == 8
    assert raw[6]["value"] is None
    assert raw[0]["unit"] == "Ct"
    summary = [m for m in data["final_results"]["measurements"] if m["replication"] is None]
    assert summary[1]["sd"] == 0.2


def test_annx_csv_and_split_export_are_equivalent(tmp_path):
    data = parse_excel_annx(str(make_workbook(tmp_path / "annx.xlsx", "ANNX")))
    raw = data["replications"][0]["records"][0]
    processed = data["processed_data"]["blocks"][0]["records"][0]
    assert raw == processed
    assert raw["count"] == 200
    assert raw["percent_gated"] == 2.74
    assert data["replications"][0]["measurements"][0]["value"] == 2.74


def test_optional_sheets_and_missing_formula_cache(tmp_path):
    import openpyxl
    path = make_workbook(tmp_path / "optional.xlsx", statistics=False)
    wb = openpyxl.load_workbook(path)
    del wb["Final results"]
    wb["Processed data_WP3_PCR_1aR1"]["B3"] = "=1+1"
    wb.save(path)
    wb.close()
    data = parse_excel_pcr(str(path))
    assert "statistical_analysis" not in data
    assert data["final_results"] == {}
    assert any("Missing cached formula" in warning for warning in data["test_details"]["parser_warnings"])
    source = data["processed_data"]["blocks"][0]["source"]
    formula = next(cell for row in source["rows"] for cell in row["cells"] if cell["coordinate"] == "B3")
    assert formula == {"column": "B", "coordinate": "B3", "value": None, "formula": "=1+1"}


def test_wrong_assay_is_rejected(tmp_path):
    with pytest.raises(ValueError, match="Expected ANNX workbook"):
        parse_excel_annx(str(make_workbook(tmp_path / "wrong.xlsx")))


def test_missing_raw_data_is_rejected(tmp_path):
    import openpyxl
    path = make_workbook(tmp_path / "no-raw.xlsx")
    wb = openpyxl.load_workbook(path)
    del wb["Raw_data_WP3_PCR_1aR1"]
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match="No PCR raw/output"):
        parse_excel_pcr(str(path))


def test_processed_final_disagreement_is_flagged_not_overwritten(tmp_path):
    import openpyxl
    path = make_workbook(tmp_path / "inconsistent.xlsx")
    wb = openpyxl.load_workbook(path)
    wb["Final results"]["D2"] = 99
    wb.save(path)
    wb.close()
    data = parse_excel_pcr(str(path))
    assert any("Processed/final disagreement" in warning for warning in data["test_details"]["parser_warnings"])
    assert data["final_results"]["measurements"][1]["value"] == 99
    assert data["processed_data"]["blocks"][0]["measurements"][1]["value"] == 2